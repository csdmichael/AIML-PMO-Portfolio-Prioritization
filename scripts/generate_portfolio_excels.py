"""
Generates per-portfolio Excel files and minimal Project Intake Template.
Outputs to data/portfolios/.
Highlights Mandatory (* [MANDATORY]) and Optional ([OPTIONAL]) columns with distinct styling.
"""

import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MANDATORY_KEYS = {
    "projectId",
    "projectName",
    "portfolio",
    "businessUnit",
    "primarySkill",
    "costEstimateUsd",
    "businessValue",
}

TEMPLATE_FIELDS = [
    ("projectName", "Project Name * [MANDATORY]", True, "Title of project / initiative (min 3 chars)"),
    ("portfolio", "Portfolio * [MANDATORY]", True, "Allowed: IT InfoSec, IT Software Development, IT Infrastructure, AI Team, Data and Analytics"),
    ("businessUnit", "Business Unit * [MANDATORY]", True, "Allowed: Semiconductor Process Equipment, Fab Operations, Global Supply Chain, Digital Transformation, etc."),
    ("primarySkill", "Primary Skill * [MANDATORY]", True, "Allowed: Application Development, Cloud and Infrastructure, Data Engineering, Cybersecurity, Network Engineering, etc."),
    ("costEstimateUsd", "Cost Estimate (USD) * [MANDATORY]", True, "Total estimated budget in USD (e.g. 1500000)"),
    ("businessValue", "Business Value (1-5) * [MANDATORY]", True, "Integer from 1 (lowest) to 5 (highest strategic impact)"),
    ("site", "Site [OPTIONAL]", False, "Allowed: Fremont, Tualatin, Singapore, Penang, Hsinchu, Global Remote (Default: Global Remote)"),
    ("executiveSponsor", "Executive Sponsor [OPTIONAL]", False, "Accountable executive sponsor name and title"),
    ("goLiveDate", "Target Go-Live Date [OPTIONAL]", False, "ISO Date YYYY-MM-DD (e.g. 2028-06-30; Default: +24 months)"),
    ("regulatorySecurity", "Regulatory / Security [OPTIONAL]", False, "TRUE if statutory, compliance, cleanroom, or cyber mandate; else FALSE"),
]

TEMPLATE_SAMPLES = [
    {
        "projectName": "Edge Wafer Defect Detection Vision System",
        "portfolio": "AI Team",
        "businessUnit": "Fab Operations",
        "primarySkill": "Data Engineering",
        "costEstimateUsd": 1850000,
        "businessValue": 5,
        "site": "Fremont",
        "executiveSponsor": "Dr. Lin Chen, VP Fab Ops",
        "goLiveDate": "2028-09-30",
        "regulatorySecurity": True,
    },
    {
        "projectName": "Zero Trust Microsegmentation Phase 2",
        "portfolio": "IT InfoSec",
        "businessUnit": "Information Security",
        "primarySkill": "Cybersecurity",
        "costEstimateUsd": 2400000,
        "businessValue": 4,
        "site": "Singapore",
        "executiveSponsor": "Dmitri Nakamura, CISO",
        "goLiveDate": "2028-04-15",
        "regulatorySecurity": True,
    },
    {
        "projectName": "Supplier Direct EDI Portal Modernization",
        "portfolio": "IT Software Development",
        "businessUnit": "Global Supply Chain",
        "primarySkill": "Application Development",
        "costEstimateUsd": 950000,
        "businessValue": 3,
        "site": "Tualatin",
        "executiveSponsor": "Elena Rostova, Dir Supply",
        "goLiveDate": "2027-11-30",
        "regulatorySecurity": False,
    },
]

def main():
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    sample_data_file = os.path.join(root_dir, "ui", "src", "app", "data", "sample-portfolio.data.ts")

    with open(sample_data_file, "r", encoding="utf-8") as f:
        text = f.read()

    idx_start = text.find('"projects": [')
    idx_end = text.find('],\n  "capacityBySkill":')
    if idx_start == -1 or idx_end == -1:
        raise ValueError("Could not locate projects array in sample-portfolio.data.ts")

    projects_json = text[idx_start + len('"projects": '):idx_end + 1]
    projects = json.loads(projects_json)

    out_dirs = [
        os.path.join(root_dir, "data", "portfolios"),
    ]
    for d in out_dirs:
        os.makedirs(d, exist_ok=True)

    headers = [
        ("projectId", "Project ID * [MANDATORY]", True),
        ("projectName", "Project Name * [MANDATORY]", True),
        ("portfolio", "Portfolio * [MANDATORY]", True),
        ("businessUnit", "Business Unit * [MANDATORY]", True),
        ("site", "Site [OPTIONAL]", False),
        ("projectManager", "Project Manager [OPTIONAL]", False),
        ("executiveSponsor", "Executive Sponsor [OPTIONAL]", False),
        ("sponsorLevel", "Sponsor Level [OPTIONAL]", False),
        ("projectStatus", "Project Status [OPTIONAL]", False),
        ("goLiveDate", "Go-Live Date [OPTIONAL]", False),
        ("monthsToGoLive", "Months to Go-Live [OPTIONAL]", False),
        ("businessValue", "Business Value (1-5) * [MANDATORY]", True),
        ("regulatorySecurity", "Regulatory / Security [OPTIONAL]", False),
        ("dtAlignment", "Digital Transformation [OPTIONAL]", False),
        ("protectedProgram", "Protected Program [OPTIONAL]", False),
        ("primarySkill", "Primary Skill * [MANDATORY]", True),
        ("skillDemandFte", "Skill Demand (FTE) [OPTIONAL]", False),
        ("costEstimateUsd", "Cost Estimate (USD) * [MANDATORY]", True),
        ("dependencyCount", "Dependency Count [OPTIONAL]", False),
        ("riskRating", "Risk Rating [OPTIONAL]", False),
        ("effortPersonWeeks", "Effort (Person-Weeks) [OPTIONAL]", False),
        ("completeness", "Completeness [OPTIONAL]", False),
        ("quality", "Quality [OPTIONAL]", False),
        ("modelCertainty", "Model Certainty [OPTIONAL]", False),
        ("confidenceScore", "Confidence Score [OPTIONAL]", False),
        ("confidenceLevel", "Confidence Level [OPTIONAL]", False),
        ("mlTier", "ML Tier [OPTIONAL]", False),
        ("finalTier", "Final Tier [OPTIONAL]", False),
        ("priorityScore", "Priority Score [OPTIONAL]", False),
        ("recommendedAction", "Recommended Action [OPTIONAL]", False),
        ("cluster", "Cluster [OPTIONAL]", False),
        ("ruleOverrides", "Rule Overrides [OPTIONAL]", False),
        ("topDrivers", "Top Drivers [OPTIONAL]", False),
        ("dataIssues", "Data Issues [OPTIONAL]", False),
        ("sourcesPresent", "Sources Present [OPTIONAL]", False),
    ]

    mandatory_fill = PatternFill(start_color="004578", end_color="004578", fill_type="solid")
    optional_fill = PatternFill(start_color="5C6873", end_color="5C6873", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    guidance_font = Font(name="Segoe UI", size=9, italic=True, color="555555")
    guidance_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    border_thin = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    def make_excel(filename, project_list, title):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title[:31]

        for col_idx, (_, label, is_mand) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = mandatory_fill if is_mand else optional_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[1].height = 28

        for row_idx, p in enumerate(project_list, 2):
            for col_idx, (k, _, _) in enumerate(headers, 1):
                val = p.get(k)
                if k == "ruleOverrides" and isinstance(val, list):
                    val = "; ".join(str(x) for x in val)
                elif k == "dataIssues" and isinstance(val, list):
                    val = "; ".join(str(x) for x in val)
                elif k == "topDrivers" and isinstance(val, list):
                    val = "; ".join(
                        f"{d.get('signalLabel')} ({d.get('shapValue'):+.2f})"
                        for d in val
                        if isinstance(d, dict)
                    )

                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = border_thin
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                elif isinstance(val, bool):
                    cell.alignment = Alignment(horizontal="center")

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 40), 14)

        for d in out_dirs:
            target_path = os.path.join(d, filename)
            wb.save(target_path)
        print(f"Generated {filename} ({len(project_list)} projects)")

    # Generate Minimal Template
    def make_template(filename):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Project Intake Template"

        # Row 1: Headers
        for col_idx, (_, label, is_mand, _) in enumerate(TEMPLATE_FIELDS, 1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = mandatory_fill if is_mand else optional_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[1].height = 32

        # Row 2: Guidance
        for col_idx, (_, _, _, guidance) in enumerate(TEMPLATE_FIELDS, 1):
            cell = ws.cell(row=2, column=col_idx, value=guidance)
            cell.fill = guidance_fill
            cell.font = guidance_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.border = border_thin

        ws.row_dimensions[2].height = 26

        # Row 3-5: Sample starter rows
        for row_idx, sample in enumerate(TEMPLATE_SAMPLES, 3):
            for col_idx, (k, _, _, _) in enumerate(TEMPLATE_FIELDS, 1):
                val = sample.get(k, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = border_thin
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                elif isinstance(val, bool):
                    cell.alignment = Alignment(horizontal="center")

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 4, 38), 16)

        for d in out_dirs:
            target_path = os.path.join(d, filename)
            wb.save(target_path)
        print(f"Generated {filename} (Minimal Template, 10 fields max)")

    # Normalize IT InfoSec to IT - InfoSec
    for p in projects:
        if p.get("portfolio") == "IT InfoSec":
            p["portfolio"] = "IT - InfoSec"

    by_portfolio = {}
    for p in projects:
        pf = p["portfolio"]
        by_portfolio.setdefault(pf, []).append(p)

    file_map = {
        "IT - InfoSec": "portfolio_it_infosec.xlsx",
        "IT InfoSec": "portfolio_it_infosec.xlsx",
        "IT Software Development": "portfolio_it_software_development.xlsx",
        "IT Infrastructure": "portfolio_it_infrastructure.xlsx",
        "AI Team": "portfolio_ai_team.xlsx",
        "Data and Analytics": "portfolio_data_and_analytics.xlsx",
    }

    import xlwt

    def make_xls(filepath, project_list, title, h_list):
        wb = xlwt.Workbook()
        ws = wb.add_sheet(title[:31])
        header_mand_style = xlwt.easyxf(
            "font: bold 1, color white, name Segoe UI; pattern: pattern solid, fore_colour dark_blue; align: horiz center, vert center"
        )
        header_opt_style = xlwt.easyxf(
            "font: bold 1, color white, name Segoe UI; pattern: pattern solid, fore_colour gray50; align: horiz center, vert center"
        )
        data_style = xlwt.easyxf("font: name Segoe UI; borders: left thin, right thin, top thin, bottom thin")
        data_num_style = xlwt.easyxf("font: name Segoe UI; align: horiz right; borders: left thin, right thin, top thin, bottom thin")
        data_bool_style = xlwt.easyxf("font: name Segoe UI; align: horiz center; borders: left thin, right thin, top thin, bottom thin")

        for col_idx, (_, label, is_mand) in enumerate(h_list):
            style = header_mand_style if is_mand else header_opt_style
            ws.write(0, col_idx, label, style)
            ws.col(col_idx).width = 256 * max(min(len(label) + 4, 38), 14)

        for row_idx, p in enumerate(project_list, 1):
            for col_idx, (k, _, _) in enumerate(h_list):
                val = p.get(k)
                if k == "ruleOverrides" and isinstance(val, list):
                    val = "; ".join(str(x) for x in val)
                elif k == "dataIssues" and isinstance(val, list):
                    val = "; ".join(str(x) for x in val)
                elif k == "topDrivers" and isinstance(val, list):
                    val = "; ".join(
                        f"{d.get('signalLabel')} ({d.get('shapValue'):+.2f})"
                        for d in val
                        if isinstance(d, dict)
                    )

                if isinstance(val, (int, float)):
                    ws.write(row_idx, col_idx, val, data_num_style)
                elif isinstance(val, bool):
                    ws.write(row_idx, col_idx, str(val).upper(), data_bool_style)
                else:
                    ws.write(row_idx, col_idx, str(val or ""), data_style)

        wb.save(filepath)

    for pf_name, p_list in by_portfolio.items():
        fname = file_map.get(pf_name, f"portfolio_{pf_name.lower().replace(' ', '_')}.xlsx")
        make_excel(fname, p_list, pf_name)

    # Specific IT - InfoSec files
    infosec_projects = by_portfolio.get("IT - InfoSec", [])
    if infosec_projects:
        make_excel("portfolio_it_-_infosec.xlsx", infosec_projects, "IT - InfoSec")
        make_excel("it_infosec_sample_projects.xlsx", infosec_projects, "IT - InfoSec")
        
        # Save directly in data/ folder as requested: "Add IT - InfoSec Portfolio and sample projects xls in data folder"
        make_excel("../it_infosec_sample_projects.xlsx", infosec_projects, "IT - InfoSec")
        make_xls(os.path.join(root_dir, "data", "it_infosec_sample_projects.xls"), infosec_projects, "IT - InfoSec", headers)
        make_xls(os.path.join(root_dir, "data", "portfolios", "it_infosec_sample_projects.xls"), infosec_projects, "IT - InfoSec", headers)
        print("Generated it_infosec_sample_projects.xls and .xlsx in data/ and data/portfolios/")

    make_excel("portfolio_all_100_projects.xlsx", projects, "All Portfolios")
    make_template("project_intake_template.xlsx")
    print("Done generating portfolio Excel files and template.")

if __name__ == "__main__":
    main()
