"""Quick add: a draft from a plain-English description and optional documents, AI prefill per field, human edits,
and submit. Every draft and every submitted project is stored only in OneLake. Rules and limits come from config.
"""

import io
import json
import uuid
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Body, Depends, File, Query, UploadFile
from pydantic import BaseModel, ConfigDict

from app.agent import clean_web_urls, is_web_url
from app.agent_context import agent_options, render_lines
from app.container import Services
from app.files import utc_now
from app.problems import FieldIssue, ProblemError
from app.routes.deps import get_services

router = APIRouter(tags=["intake"])
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json"}


class DraftCreate(BaseModel):
    projectName: str
    description: str


class DraftDocument(BaseModel):
    uploadId: str
    fileName: str
    sizeBytes: int
    oneLakePath: str
    uploadedAtUtc: str
    textExtraction: str


class DraftView(BaseModel):
    draftId: str
    status: str
    projectName: str
    description: str
    documents: list[DraftDocument]
    fields: dict[str, Any]
    missingRequired: list[str]
    model: dict[str, Any] | None = None
    createdAtUtc: str
    updatedAtUtc: str
    projectId: str | None = None


class SubmitResult(BaseModel):
    projectId: str
    status: str
    pipelineRunId: str | None
    warnings: list[str]


def _draft_view(draft: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in draft.items() if key != "documentText"}


def _load_draft(services: Services, draft_id: str) -> tuple[dict[str, Any], str]:
    path = services.paths.draft(draft_id)
    found = services.entries.get_with_etag(path)
    if found is None:
        raise ProblemError(404, f"Draft {draft_id} was not found. Expected it at {path} in OneLake.", slug="not-found")
    return found


def _save_draft(services: Services, draft_id: str, draft: dict[str, Any], etag: str) -> None:
    draft["updatedAtUtc"] = utc_now()
    services.entries.replace(services.paths.draft(draft_id), draft, etag)


def _missing_required(services: Services, fields: dict[str, Any]) -> list[str]:
    missing = []
    for rule in services.config.data_entry.projectFields:
        value = fields.get(rule.key, {}).get("value")
        if rule.required and (value is None or value == "" or value == []):
            missing.append(rule.key)
    return missing


def _extract_text(extension: str, data: bytes, max_chars: int, xlsx_rows: int) -> tuple[str, str]:
    if extension in TEXT_EXTENSIONS:
        return data.decode("utf-8", errors="replace")[:max_chars], "utf8"
    if extension == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_number > xlsx_rows:
                    break
                lines.append(",".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(lines)[:max_chars], "xlsx-rows"
    return "", "stored-only"


def _typed(rule: Any, raw: Any) -> Any:
    if isinstance(raw, str):
        text = raw.strip()
        if text == "":
            return None
        if rule.type in ("integer", "number"):
            try:
                return float(text.replace(",", "").replace("$", ""))
            except ValueError:
                return raw
        if rule.type == "boolean" and text.lower() in ("true", "false"):
            return text.lower() == "true"
    return raw


def _field_issues(issues: list[FieldIssue]) -> dict[str, str]:
    messages: dict[str, list[str]] = {}
    for issue in issues:
        key = issue.field.split("[")[0].split(".")[0]
        messages.setdefault(key, []).append(issue.message)
    return {key: "; ".join(items) for key, items in messages.items()}


@router.post("/intake/drafts", status_code=201, response_model=DraftView, operation_id="createIntakeDraft", summary="Start a quick-add draft")
def create_draft(body: DraftCreate, services: Services = Depends(get_services)) -> dict[str, Any]:
    limits = services.config.ai_prefill.limits
    name = body.projectName.strip()
    description = body.description.strip()
    if not 1 <= len(name) <= limits.maxProjectNameChars:
        raise ProblemError(422, "The project name is not valid.", slug="validation-failed", errors=[FieldIssue("projectName", f"must be 1 to {limits.maxProjectNameChars} characters")])
    if not limits.minDescriptionChars <= len(description) <= limits.maxDescriptionChars:
        raise ProblemError(422, "The description is not valid.", slug="validation-failed", errors=[FieldIssue("description", f"must be {limits.minDescriptionChars} to {limits.maxDescriptionChars} characters")])
    draft_id = f"DRF-{uuid.uuid4().hex[:12].upper()}"
    now = utc_now()
    draft = {"draftId": draft_id, "status": "draft", "projectName": name, "description": description, "documents": [], "fields": {}, "missingRequired": [], "model": None, "createdAtUtc": now, "updatedAtUtc": now, "projectId": None}
    services.entries.create(services.paths.draft(draft_id), draft)
    return _draft_view(draft)


@router.post("/intake/drafts/{draftId}/documents", status_code=201, response_model=DraftDocument, operation_id="addIntakeDocument", summary="Attach a supporting document to a draft")
def add_document(draftId: str, file: UploadFile = File(...), services: Services = Depends(get_services)) -> dict[str, Any]:
    limits = services.config.ai_prefill.limits
    draft, etag = _load_draft(services, draftId)
    if draft["status"] != "draft":
        raise ProblemError(409, "Documents can only be added before the prefill runs.", slug="conflict")
    if len(draft["documents"]) >= limits.maxDocuments:
        raise ProblemError(422, f"A draft accepts at most {limits.maxDocuments} documents.", slug="validation-failed")
    name = Path(file.filename or "document").name
    extension = Path(name).suffix.lower()
    if extension not in limits.acceptedDocumentExtensions:
        raise ProblemError(422, f"{extension or 'This file type'} is not accepted.", slug="validation-failed", errors=[FieldIssue("file", "unsupported type")])
    data = file.file.read(limits.maxDocumentBytes + 1)
    if len(data) > limits.maxDocumentBytes:
        raise ProblemError(413, "The document is larger than the allowed size.", slug="payload-too-large")
    upload_id = uuid.uuid4().hex[:10]
    path = services.paths.named_file(services.paths.draft_documents_folder(draftId), upload_id, name)
    services.lake.write_bytes(path, data)
    text, extraction = _extract_text(extension, data, limits.maxDocumentTextChars, limits.xlsxRowsPerSheet)
    meta = {"uploadId": upload_id, "fileName": name, "sizeBytes": len(data), "oneLakePath": path, "uploadedAtUtc": utc_now(), "textExtraction": extraction}
    draft["documents"].append({**meta, "text": text})
    _save_draft(services, draftId, draft, etag)
    return meta


@router.get("/intake/drafts/{draftId}", response_model=DraftView, operation_id="getIntakeDraft", summary="A draft with its fields and provenance")
def get_draft(draftId: str, services: Services = Depends(get_services)) -> dict[str, Any]:
    draft, _ = _load_draft(services, draftId)
    return _draft_view(draft)


class PrefillFieldAnswer(BaseModel):
    model_config = ConfigDict(extra="ignore")
    value: Any = None
    evidence: str = "notProvided"
    evidenceQuote: str | None = None
    rationale: str = ""
    sourceUrl: str | None = None


class PrefillDraft(BaseModel):
    """The intake agent's reply: one entry per field key, plus the public URLs it used."""

    model_config = ConfigDict(extra="ignore")
    fields: dict[str, PrefillFieldAnswer]
    sourceUrls: list[str] = []


@router.post("/intake/drafts/{draftId}/prefill", response_model=DraftView, operation_id="prefillIntakeDraft", summary="AI prefill of every project field with evidence")
def prefill(draftId: str, services: Services = Depends(get_services)) -> dict[str, Any]:
    draft, etag = _load_draft(services, draftId)
    if draft["status"] not in ("draft", "prefilled"):
        raise ProblemError(409, "This draft has already been submitted.", slug="conflict")
    prefill_settings = services.config.ai_prefill
    contract = services.config.data_entry
    documents = "\n\n".join(f"[{doc['fileName']}]\n{doc.get('text', '')}" for doc in draft["documents"]) or "(none)"
    guidance = {item.key: item.guidance for item in prefill_settings.fields}
    options = {**agent_options(services), "fieldGuidance": guidance}
    user_message = render_lines(
        prefill_settings.prompt.userTemplate,
        projectName=draft["projectName"],
        description=draft["description"],
        documents=documents,
        options=json.dumps(options, ensure_ascii=False),
    )
    proposal, run = services.intake_agent.structured(
        [
            {"role": "system", "content": "\n".join(prefill_settings.prompt.system)},
            {"role": "user", "content": user_message},
        ],
        PrefillDraft,
    )
    proposals = {rule.key: proposal.fields.get(rule.key, PrefillFieldAnswer()) for rule in contract.projectFields}
    proposals["projectName"] = PrefillFieldAnswer(
        value=draft["projectName"], evidence="explicit", evidenceQuote=draft["projectName"], rationale="Entered by the user.",
    )
    typed: dict[str, Any] = {}
    for rule in contract.projectFields:
        value = _typed(rule, proposals[rule.key].value)
        if value is not None:
            typed[rule.key] = value
    checked = services.projects.validate(typed, partial=True)
    problems = _field_issues(checked.issues)
    previous = draft.get("fields") or {}
    fields: dict[str, Any] = {}
    for rule in contract.projectFields:
        kept = previous.get(rule.key) or {}
        if kept.get("source") == "human" and kept.get("value") not in (None, "", []):
            fields[rule.key] = kept
            continue
        item = proposals[rule.key]
        value = checked.values.get(rule.key)
        evidence = item.evidence if value is not None and item.evidence in prefill_settings.evidenceToConfidence else "notProvided"
        mapping = prefill_settings.evidenceToConfidence[evidence]
        rationale = item.rationale
        if rule.key in problems:
            rationale = f"{rationale} (Not used: {problems[rule.key]})".strip()
        fields[rule.key] = {
            "value": value, "aiValue": value, "source": "ai", "evidence": evidence,
            "confidenceLevel": mapping.level, "confidenceScore": mapping.score,
            "evidenceQuote": item.evidenceQuote, "rationale": rationale,
            "sourceUrl": item.sourceUrl if item.sourceUrl and is_web_url(item.sourceUrl) else None,
            "updatedAtUtc": utc_now(),
        }
    agent = services.intake_agent.settings
    draft.update(
        status="prefilled",
        fields=fields,
        missingRequired=_missing_required(services, fields),
        model={
            "agentName": agent.agentName, "agentVersion": agent.agentVersion,
            "responseId": run.response_id, "queryIds": list(run.query_ids), "webSearches": run.web_searches,
            "sourceUrls": clean_web_urls(proposal.sourceUrls), "requestedAtUtc": utc_now(),
        },
    )
    _save_draft(services, draftId, draft, etag)
    return _draft_view(draft)


@router.patch("/intake/drafts/{draftId}/fields", response_model=DraftView, operation_id="editIntakeFields", summary="Human edits to prefilled fields")
def edit_fields(draftId: str, body: dict[str, Any] = Body(...), services: Services = Depends(get_services)) -> dict[str, Any]:
    draft, etag = _load_draft(services, draftId)
    if draft["status"] == "submitted":
        raise ProblemError(409, "This draft has already been submitted.", slug="conflict")
    known = {rule.key: rule for rule in services.config.data_entry.projectFields}
    unknown = [FieldIssue(key, "is not a project field") for key in body if key not in known]
    if unknown:
        raise ProblemError(422, "Some fields are not valid.", slug="validation-failed", errors=unknown)
    edits = {key: _typed(known[key], raw) for key, raw in body.items()}
    current = {key: field.get("value") for key, field in draft["fields"].items()}
    merged = {key: value for key, value in {**current, **edits}.items() if value is not None}
    checked = services.projects.validate(merged, partial=True)
    problems = _field_issues(checked.issues)
    if problems:
        raise ProblemError(422, "Some fields are not valid.", slug="validation-failed", errors=[FieldIssue(key, message) for key, message in problems.items()])
    for key in edits:
        previous = draft["fields"].get(key, {})
        draft["fields"][key] = {**previous, "value": checked.values.get(key), "source": "human", "evidence": "human", "confidenceLevel": "Human", "confidenceScore": None, "updatedAtUtc": utc_now()}
    draft["missingRequired"] = _missing_required(services, draft["fields"])
    _save_draft(services, draftId, draft, etag)
    return _draft_view(draft)


@router.post("/intake/drafts/{draftId}/submit", status_code=201, response_model=SubmitResult, operation_id="submitIntakeDraft", summary="Create the project from the reviewed draft")
def submit(draftId: str, runPipeline: bool = Query(True), services: Services = Depends(get_services)) -> dict[str, Any]:
    draft, etag = _load_draft(services, draftId)
    if draft["status"] == "submitted":
        raise ProblemError(409, "This draft has already been submitted.", slug="conflict")
    values = {key: field["value"] for key, field in draft["fields"].items() if field.get("value") is not None}
    values.setdefault("projectName", draft["projectName"])
    provenance = {key: {k: v for k, v in field.items() if k != "value"} for key, field in draft["fields"].items()}
    documents = [{"fileName": d["fileName"], "oneLakePath": d["oneLakePath"]} for d in draft["documents"]]
    entry = services.projects.create(values, extra={"fieldProvenance": provenance, "intakeDraftId": draftId, "documents": documents})
    run_id, warnings = services.pipeline.start_optional() if runPipeline else (None, [])
    draft.update(status="submitted", projectId=entry["projectId"], submittedAtUtc=utc_now())
    _save_draft(services, draftId, draft, etag)
    return {"projectId": entry["projectId"], "status": services.config.api.projects.statusPending, "pipelineRunId": run_id, "warnings": warnings}
