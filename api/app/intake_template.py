"""Builds the intake workbook template from config: a header row and dropdown lists, with no data rows."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.config import AppConfig, ConfigError
from app.contracts import static_option_values

def _intake_table(config: AppConfig) -> tuple[str, int]:
    workbook = config.api.templates.intakeWorkbook
    for source in config.data_sources.sources:
        if source.key != workbook.sourceKey:
            continue
        for table in source.tables:
            if table.name == workbook.tableName:
                if not table.sheet:
                    raise ConfigError(f"data-sources.json table '{table.name}' has no sheet")
                return table.sheet, table.headerRow or 1
    raise ConfigError(f"data-sources.json has no table '{workbook.tableName}' in source '{workbook.sourceKey}'")


def build_intake_template(config: AppConfig) -> bytes:
    settings = config.api.templates.intakeWorkbook
    sheet_name, header_row = _intake_table(config)
    headers = [column.header for column in settings.identifierColumns] + settings.signalColumns

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for _ in range(header_row - 1):
        sheet.append([])
    sheet.append(headers)
    sheet.freeze_panes = f"A{header_row + 1}"

    lists = workbook.create_sheet(settings.listsSheetName)
    rules = {rule.key: rule for rule in config.data_entry.projectFields}
    list_column = 0
    for position, column in enumerate(settings.identifierColumns, start=1):
        if column.fieldKey is None or column.fieldKey not in rules:
            continue
        rule = rules[column.fieldKey]
        values = static_option_values(config.data_entry, rule) or []
        if not values:
            continue
        list_column += 1
        list_letter = get_column_letter(list_column)
        lists.cell(row=1, column=list_column, value=rule.label)
        for row, value in enumerate(values, start=2):
            lists.cell(row=row, column=list_column, value=value)
        formula = f"'{settings.listsSheetName}'!${list_letter}$2:${list_letter}${len(values) + 1}"
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        sheet.add_data_validation(validation)
        target = get_column_letter(position)
        validation.add(f"{target}{header_row + 1}:{target}{header_row + settings.dropdownRows}")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
