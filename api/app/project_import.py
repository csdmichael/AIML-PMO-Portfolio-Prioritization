"""Bulk project import from an Excel register: one project per row, with every row checked before any project is saved."""

import datetime as dt
import io
from typing import Any

import openpyxl
from openpyxl import Workbook

from app.catalog import Catalog, CatalogOptions
from app.config import ProjectImportSettings
from app.contracts import DataEntryContract, FieldRule, static_option_values
from app.problems import FieldIssue, ProblemError
from app.projects import ProjectService

ALLOWED_VALUES_SHEET = "Allowed values"


def _normal(text: object) -> str:
    return " ".join(str(text).split()).casefold()


class ProjectImport:
    def __init__(
        self,
        *,
        settings: ProjectImportSettings,
        contract: DataEntryContract,
        catalog: Catalog,
        projects: ProjectService,
    ) -> None:
        self._settings = settings
        self._contract = contract
        self._catalog = catalog
        self._projects = projects
        self._rules = {rule.key: rule for rule in contract.projectFields}
        self._true_words = {_normal(word) for word in settings.booleanWords["true"]}
        self._false_words = {_normal(word) for word in settings.booleanWords["false"]}

    def import_workbook(self, data: bytes) -> list[dict[str, Any]]:
        rows = self._rows(data)
        plans = self._plan(rows)
        created: list[dict[str, Any]] = []
        for values in plans:
            entry = self._projects.write_entry(values)
            created.append({"projectId": entry["projectId"], "projectName": entry.get("projectName")})
        return created

    def template(self) -> bytes:
        """The import workbook: the header row, and the allowed values for the columns that have a list."""
        workbook = Workbook()
        sheet = workbook.active
        if self._settings.sheetName:
            sheet.title = self._settings.sheetName
        for _ in range(self._settings.headerRow - 1):
            sheet.append([])
        sheet.append([column.header for column in self._settings.columns])
        sheet.freeze_panes = f"A{self._settings.headerRow + 1}"
        allowed = workbook.create_sheet(ALLOWED_VALUES_SHEET)
        column_number = 0
        for column in self._settings.columns:
            values = self._allowed_for_template(column)
            if not values:
                continue
            column_number += 1
            allowed.cell(row=1, column=column_number, value=column.header)
            for row_number, value in enumerate(values, start=2):
                allowed.cell(row=row_number, column=column_number, value=value)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _allowed_for_template(self, column: Any) -> list[str] | None:
        if column.lookup == "portfolioName":
            return sorted({str(item["displayName"]) for item in self._catalog.portfolios() if item.get("displayName")})
        if column.lookup == "programName":
            return sorted({str(program["displayName"]) for item in self._catalog.portfolios() for program in item.get("programs") or [] if program.get("displayName")})
        if column.lookup == "skillName":
            return [skill.displayName for skill in self._contract.options.skills]
        rule = self._rules[column.fieldKey]
        return static_option_values(self._contract, rule) if rule.type == "select" else None

    def _rows(self, data: bytes) -> list[tuple[int, dict[str, Any]]]:
        if not data:
            raise ProblemError(422, "The file is empty.", slug="validation-failed", errors=[FieldIssue("file", "is empty")])
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as error:
            raise ProblemError(
                422, "The file is not a readable Excel workbook.", slug="validation-failed", errors=[FieldIssue("file", "must be an Excel .xlsx workbook")],
            ) from error
        try:
            if self._settings.sheetName:
                if self._settings.sheetName not in workbook.sheetnames:
                    raise ProblemError(422, "The workbook has no sheet for projects.", slug="validation-failed", errors=[FieldIssue("file", f"must have a sheet named {self._settings.sheetName}")])
                sheet = workbook[self._settings.sheetName]
            else:
                sheet = workbook.worksheets[0]
            cells = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        header_index = self._settings.headerRow - 1
        header = cells[header_index] if len(cells) > header_index else ()
        by_header = {_normal(column.header): column for column in self._settings.columns}
        positions: dict[str, int] = {}
        for index, title in enumerate(header):
            column = by_header.get(_normal(title)) if title is not None else None
            if column is not None:
                positions[column.fieldKey] = index
        missing = [
            FieldIssue(column.header, "is missing from the header row")
            for column in self._settings.columns
            if column.required and column.fieldKey not in positions
        ]
        if missing:
            raise ProblemError(422, "The header row is missing required columns.", slug="validation-failed", errors=missing)
        rows: list[tuple[int, dict[str, Any]]] = []
        for number, row in enumerate(cells[header_index + 1 :], start=self._settings.headerRow + 1):
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in row):
                continue
            if len(rows) >= self._settings.maxRows:
                raise ProblemError(
                    422, "The workbook has more project rows than the import accepts.", slug="validation-failed",
                    errors=[FieldIssue("file", f"must have at most {self._settings.maxRows} project rows")],
                )
            rows.append((number, {key: row[index] if index < len(row) else None for key, index in positions.items()}))
        if not rows:
            raise ProblemError(422, "The workbook has no project rows.", slug="validation-failed", errors=[FieldIssue("file", "has no project rows")])
        return rows

    def _plan(self, rows: list[tuple[int, dict[str, Any]]]) -> list[dict[str, Any]]:
        headers = {column.fieldKey: column.header for column in self._settings.columns}
        lookups = {column.fieldKey: column.lookup for column in self._settings.columns if column.lookup}
        options = self._catalog.options()
        portfolios = self._catalog.portfolio_index()
        plans: list[dict[str, Any]] = []
        issues: list[FieldIssue] = []
        for number, raw in rows:
            values = {key: self._cell(self._rules[key], value) for key, value in raw.items()}
            self._resolve(values, lookups, options, portfolios)
            result = self._projects.validate(values)
            if result.valid:
                plans.append(result.values)
                continue
            for issue in result.issues:
                field = issue.field.split("[")[0].split(".")[0]
                issues.append(FieldIssue(f"row {number}, {headers.get(field, field)}", issue.message))
        if issues:
            raise ProblemError(422, "Some rows do not match the data-entry rules. No project was saved.", slug="validation-failed", errors=issues)
        return plans

    def _cell(self, rule: FieldRule, raw: Any) -> Any:
        if raw is None:
            return None
        if isinstance(raw, str):
            text = raw.strip()
            if not text:
                return None
            if rule.type == "boolean":
                word = _normal(text)
                if word in self._true_words:
                    return True
                if word in self._false_words:
                    return False
                return text
            if rule.type in ("integer", "number"):
                try:
                    return float(text.replace(",", "").replace("$", ""))
                except ValueError:
                    return text
            return text
        if rule.type == "date" and isinstance(raw, dt.datetime):
            return raw.date().isoformat()
        if rule.type == "date" and isinstance(raw, dt.date):
            return raw.isoformat()
        return raw

    def _resolve(
        self,
        values: dict[str, Any],
        lookups: dict[str, str],
        options: CatalogOptions,
        portfolios: dict[str, dict[str, Any]],
    ) -> None:
        """Replace names from the workbook with the keys the contract uses. Unknown names stay as typed and fail validation."""
        for key, lookup in lookups.items():
            value = values.get(key)
            if not isinstance(value, str):
                continue
            if lookup == "portfolioName":
                values[key] = options.key_for_name("portfolios", value) or value
            elif lookup == "programName":
                programs = (portfolios.get(str(values.get("portfolioKey") or "")) or {}).get("programs") or []
                match = next((item for item in programs if _normal(item.get("displayName", "")) == _normal(value)), None)
                values[key] = match["key"] if match else value
            elif lookup == "skillName":
                match = next((skill for skill in self._contract.options.skills if _normal(skill.displayName) == _normal(value)), None)
                values[key] = match.key if match else value
