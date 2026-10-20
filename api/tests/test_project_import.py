import datetime as dt
import io
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.cache import TtlCache
from app.catalog import Catalog
from app.config import load_config
from app.lake_paths import LakePaths
from app.main import create_app
from app.options_store import OptionsStore
from app.project_changes import ProjectChanges
from app.project_import import ProjectImport
from app.projects import ProjectService
from app.routes.deps import get_services
from app.store import EntryStore, ExportReader
from tests.lake_doubles import MemoryLake

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PORTFOLIOS = [
    {
        "key": "it-infosec",
        "displayName": "IT InfoSec",
        "code": "INFOSEC",
        "role": "main",
        "programs": [
            {"key": "zero-trust-identity", "displayName": "Zero Trust and Identity", "code": "ZTI"},
            {"key": "security-operations", "displayName": "Security Operations", "code": "SOC"},
        ],
    },
    {
        "key": "data-analytics",
        "displayName": "Data and Analytics",
        "code": "DATA",
        "role": "coupled",
        "programs": [{"key": "data-platform", "displayName": "Data Platform and Lakehouse", "code": "DPL"}],
    },
]


def valid_row(config, **overrides):
    options = config.data_entry.options
    row = {
        "Project name": "HR Policies Website",
        "Portfolio": "IT InfoSec",
        "Program": "Zero Trust and Identity",
        "Business unit": options.businessUnits[0],
        "Site": options.sites[0],
        "Executive sponsor": "Dana Park",
        "Status": options.projectStatuses[0],
        "Sponsor level": options.sponsorLevels[0],
        "Go-live date": dt.date(2027, 6, 30),
        "Business value estimate (1 to 5)": 4,
        "User base (number of users)": 1500,
        "Regulatory or security driven": "No",
        "Digital transformation aligned": "Yes",
        "Primary skill": options.skills[0].displayName,
        "Skill demand (FTE)": 2,
        "Cost estimate (USD)": 1200000,
        "Risk rating": options.riskRatings[0],
        "Effort (person-weeks)": 26,
    }
    row.update(overrides)
    return row


def workbook(config, rows, *, headers=None):
    columns = headers or [column.header for column in config.api.projectImport.columns]
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(header) for header in columns])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def harness():
    config = load_config()
    lake = MemoryLake()
    paths = LakePaths(config)
    lake.files[paths.portfolio_export()] = ({"portfolios": PORTFOLIOS, "projects": []}, "etag-export")
    entries = EntryStore(lake)
    changes = ProjectChanges(entries=entries, paths=paths, contract=config.data_entry, cache=TtlCache(0))
    exports = ExportReader(lake=lake, cache=TtlCache(0), pipeline_name="PMO pipeline")
    catalog = Catalog(
        exports=exports, entries=entries, paths=paths, statuses=config.api.projects, contract=config.data_entry, changes=changes,
    )
    projects = ProjectService(
        lake=lake,
        entries=entries,
        paths=paths,
        catalog=catalog,
        contract=config.data_entry,
        validation=config.api.validation,
        project_settings=config.api.projects,
        upload_settings=config.api.uploads,
        changes=changes,
    )
    services = SimpleNamespace(
        config=config,
        paths=paths,
        catalog=catalog,
        projects=projects,
        imports=ProjectImport(settings=config.api.projectImport, contract=config.data_entry, catalog=catalog, projects=projects),
        pipeline=SimpleNamespace(start_optional=lambda: (None, [])),
        options=OptionsStore(lake=lake, contract=config.data_entry, path=config.api.oneLake.lookupsFile, ttl_seconds=0),
    )
    app = create_app()
    app.dependency_overrides[get_services] = lambda: services
    with TestClient(app) as client:
        yield SimpleNamespace(client=client, lake=lake, paths=paths, config=config)


def upload(client, data, name="register.xlsx"):
    return client.post("/v1/projects/import", files={"file": (name, data, XLSX)})


def saved_projects(harness):
    folder = harness.paths.project_entries_folder()
    return [path for path in harness.lake.files if path.startswith(folder + "/")]


def test_valid_rows_become_pending_projects_with_their_keys_and_typed_values(harness):
    data = workbook(harness.config, [valid_row(harness.config), valid_row(harness.config, **{"Project name": "Second Project"})])
    response = upload(harness.client, data)
    assert response.status_code == 201
    body = response.json()
    assert body["count"] == 2
    assert all(item["projectId"].startswith("GIS-") for item in body["created"])
    first = harness.lake.files[harness.paths.project_entry(body["created"][0]["projectId"])][0]
    assert first["portfolioKey"] == "it-infosec"
    assert first["programKey"] == "zero-trust-identity"
    assert first["primarySkill"] == harness.config.data_entry.options.skills[0].key
    assert first["goLiveDate"] == "2027-06-30"
    assert first["regulatorySecurity"] is False
    assert first["dtAlignment"] is True
    assert first["userBase"] == 1500


def test_one_bad_row_saves_nothing_and_names_the_row(harness):
    bad = valid_row(harness.config, Status="Nonsense")
    response = upload(harness.client, workbook(harness.config, [valid_row(harness.config), bad]))
    assert response.status_code == 422
    assert "row 3" in response.text
    assert "Status" in response.text
    assert saved_projects(harness) == []


def test_a_missing_required_column_is_reported_before_any_row(harness):
    headers = [column.header for column in harness.config.api.projectImport.columns if column.header != "Site"]
    response = upload(harness.client, workbook(harness.config, [valid_row(harness.config)], headers=headers))
    assert response.status_code == 422
    assert "Site" in response.text
    assert saved_projects(harness) == []


def test_a_file_that_is_not_a_workbook_is_rejected(harness):
    response = upload(harness.client, b"not an excel file", name="register.csv")
    assert response.status_code == 422
    assert saved_projects(harness) == []


def test_an_unknown_portfolio_name_is_reported_on_its_row(harness):
    response = upload(harness.client, workbook(harness.config, [valid_row(harness.config, Portfolio="No Such Portfolio")]))
    assert response.status_code == 422
    assert "row 2, Portfolio" in response.text


def test_program_names_are_matched_within_the_portfolio_of_the_same_row(harness):
    row = valid_row(harness.config, Program="Security Operations")
    response = upload(harness.client, workbook(harness.config, [row]))
    assert response.status_code == 201
    project_id = response.json()["created"][0]["projectId"]
    assert harness.lake.files[harness.paths.project_entry(project_id)][0]["programKey"] == "security-operations"


def test_a_program_from_another_portfolio_is_rejected(harness):
    row = valid_row(harness.config, Program="Data Platform and Lakehouse")
    response = upload(harness.client, workbook(harness.config, [row]))
    assert response.status_code == 422
    assert "Program" in response.text


def test_a_workbook_with_no_project_rows_is_rejected(harness):
    response = upload(harness.client, workbook(harness.config, []))
    assert response.status_code == 422
    assert "no project rows" in response.text


def test_more_rows_than_the_limit_are_refused(harness):
    limit = harness.config.api.projectImport.maxRows
    rows = [valid_row(harness.config, **{"Project name": f"Project {number}"}) for number in range(limit + 1)]
    response = upload(harness.client, workbook(harness.config, rows))
    assert response.status_code == 422
    assert saved_projects(harness) == []


def test_the_template_has_the_headers_and_the_allowed_values(harness):
    response = harness.client.get("/v1/templates/project-import.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == harness.config.api.projectImport.contentType
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    headers = [cell.value for cell in book.worksheets[0][1]]
    assert headers == [column.header for column in harness.config.api.projectImport.columns]
    allowed = {cell.value for row in book["Allowed values"].iter_rows() for cell in row}
    assert {"IT InfoSec", "Zero Trust and Identity", "Data Platform and Lakehouse"} <= allowed
    assert harness.config.data_entry.options.sites[0] in allowed
